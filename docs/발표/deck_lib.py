# -*- coding: utf-8 -*-
"""발표 PPT 초안 생성기 — python-pptx. 테마 'Graph Ink' (흰 바탕 · 남색 · 코랄 강조 · 노드-엣지 모티프)."""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.oxml.ns import qn
from lxml import etree

THEME = "A"
THEMES = {
    "A": dict(name="Graph Ink", bg="FFFFFF", text="1F2430", muted="6B7280", primary="2F3C7E", accent="F96167",
              good="1A9E7A", card="F3F4F8", card2="EEF1FB", dark="1B2140", dark_text="FFFFFF", dark_muted="B8C0E0", mono="Consolas"),
    "B": dict(name="Dark Board", bg="0F172A", text="E5E7EB", muted="94A3B8", primary="5EEAD4", accent="FB7185",
              good="34D399", card="1E293B", card2="172033", dark="0B1020", dark_text="F8FAFC", dark_muted="94A3B8", mono="Consolas"),
    "C": dict(name="Signal Red", bg="FFFFFF", text="1A1A1A", muted="6B6B6B", primary="A50034", accent="A50034",
              good="2E7D32", card="F5F5F5", card2="FBECEF", dark="2B0A14", dark_text="FFFFFF", dark_muted="D9B8C2", mono="Consolas"),
}
T = THEMES[THEME]
FONT = "Malgun Gothic"
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(W); prs.slide_height = Inches(H)
BLANK = prs.slide_layouts[6]


def rgb(h): return RGBColor.from_string(h)


def rect(slide, x, y, w, h, fill, shape=MSO_SHAPE.RECTANGLE, line=None, radius=None, shadow=False):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = rgb(fill)
    if line: s.line.color.rgb = rgb(line); s.line.width = Pt(0.75)
    else: s.line.fill.background()
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        s.adjustments[0] = radius
    if not shadow:
        # 그림자 제거
        spPr = s._element.spPr
        el = etree.SubElement(spPr, qn("a:effectLst"))
    s.text_frame.text = ""
    return s


def tb(slide, x, y, w, h, text, size=14, bold=False, color=None, align=PP_ALIGN.LEFT, font=FONT, anchor=MSO_ANCHOR.TOP,
       line_spacing=1.15, italic=False):
    """텍스트 상자. text 는 str 또는 [(text, {opts}), ...] 런 목록 또는 문단 목록(list of list)."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    paras = text if isinstance(text, list) and text and isinstance(text[0], list) else [text]
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align; p.line_spacing = line_spacing
        runs = para if isinstance(para, list) else [(para, {})]
        for rtext, opts in (runs if isinstance(runs[0], tuple) else [(runs, {})]):
            r = p.add_run(); r.text = rtext
            f = r.font; f.name = opts.get("font", font); f.size = Pt(opts.get("size", size)); f.bold = opts.get("bold", bold)
            f.italic = opts.get("italic", italic)
            f.color.rgb = rgb(opts.get("color", color or T["text"]))
    return box


def bullets(slide, x, y, w, h, items, size=14, color=None, gap=6, bullet_color=None, font=FONT):
    """글머리 기호 문단들. items: str 또는 (str, level) 또는 [(run, opts), ...]."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, it in enumerate(items):
        level = 0
        if isinstance(it, tuple) and len(it) == 2 and isinstance(it[1], int):
            it, level = it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.12; p.space_after = Pt(gap)
        pPr = p._p.get_or_add_pPr()
        indent = 0.22 + 0.25 * level
        pPr.set("marL", str(int(Inches(indent)))); pPr.set("indent", str(-int(Inches(0.22))))
        bu = etree.SubElement(pPr, qn("a:buClr")); srgb = etree.SubElement(bu, qn("a:srgbClr")); srgb.set("val", bullet_color or T["primary"])
        buf = etree.SubElement(pPr, qn("a:buFont")); buf.set("typeface", "Arial")
        bc = etree.SubElement(pPr, qn("a:buChar")); bc.set("char", "•" if level == 0 else "–")
        runs = it if isinstance(it, list) else [(it, {})]
        for rtext, opts in runs:
            r = p.add_run(); r.text = rtext
            f = r.font; f.name = opts.get("font", font); f.size = Pt(opts.get("size", size - (1 if level else 0)))
            f.bold = opts.get("bold", False); f.color.rgb = rgb(opts.get("color", color or T["text"]))
    return box


def node_motif(slide, x, y, scale=1.0, colors=None):
    """노드-엣지 모티프: 원 3개 + 선 2개."""
    colors = colors or [T["primary"], T["accent"], T["primary"]]
    pts = [(0, 0.25), (0.55, 0), (1.05, 0.3)]
    r = 0.22 * scale
    # 선 먼저
    for (ax, ay), (bx, by) in [(pts[0], pts[1]), (pts[1], pts[2])]:
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x + ax * scale + r / 2), Inches(y + ay * scale + r / 2),
                                       Inches(x + bx * scale + r / 2), Inches(y + by * scale + r / 2))
        c.line.color.rgb = rgb(T["muted"]); c.line.width = Pt(1.5 * scale)
    for (px, py), col in zip(pts, colors):
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + px * scale), Inches(y + py * scale), Inches(r), Inches(r))
        o.fill.solid(); o.fill.fore_color.rgb = rgb(col); o.line.fill.background()
        etree.SubElement(o._element.spPr, qn("a:effectLst"))


SECTION = [0]


def header(slide, num, title, subtitle=None, dark=False):
    """장 번호(연한 큰 숫자, 자동 증가) + 제목 + 부제. 모티프는 우상단."""
    if num:
        SECTION[0] += 1; num = SECTION[0]
    bg = T["dark"] if dark else T["bg"]
    rect(slide, 0, 0, W, H, bg)
    col_t = T["dark_text"] if dark else T["text"]
    col_m = T["dark_muted"] if dark else T["muted"]
    if num:
        tb(slide, 0.6, 0.42, 1.4, 1.0, f"{num:02d}", size=40, bold=True, color=(T["card2"] if not dark else "2A3160") if THEME != "B" else "1E293B", font="Arial")
    tb(slide, 0.6 if not num else 1.55, 0.55, 9.6, 0.7, title, size=28, bold=True, color=col_t)
    if subtitle:
        tb(slide, 0.6 if not num else 1.55, 1.18, 10.3, 0.5, subtitle, size=14, color=col_m)
    node_motif(slide, 11.55, 0.55, scale=0.9)
    # 페이지 번호
    tb(slide, 12.2, 7.05, 0.8, 0.3, str(len(prs.slides)), size=10, color=col_m, align=PP_ALIGN.RIGHT, font="Arial")
    tb(slide, 0.6, 7.05, 6, 0.3, "역량 경로 추천 — 전공 ↔ 역량 ↔ LG 계열사 신입 직무 지식그래프", size=9, color=col_m)


def card(slide, x, y, w, h, title=None, body=None, fill=None, title_color=None, body_size=12, title_size=14, icon=None):
    rect(slide, x, y, w, h, fill or T["card"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
    cy = y + 0.18
    if icon:
        o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.22), Inches(cy - 0.02), Inches(0.42), Inches(0.42))
        o.fill.solid(); o.fill.fore_color.rgb = rgb(title_color or T["primary"]); o.line.fill.background()
        etree.SubElement(o._element.spPr, qn("a:effectLst"))
        tf = o.text_frame; tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER; r = p.add_run(); r.text = icon
        r.font.size = Pt(12); r.font.bold = True; r.font.color.rgb = rgb("FFFFFF"); r.font.name = "Arial"
    if title:
        tb(slide, x + (0.78 if icon else 0.25), cy + 0.03, w - (1.0 if icon else 0.5), 0.4, title, size=title_size, bold=True, color=title_color or T["primary"])
        cy += 0.5
    if body:
        if isinstance(body, list):
            bullets(slide, x + 0.25, cy + 0.05, w - 0.5, h - (cy - y) - 0.2, body, size=body_size)
        else:
            tb(slide, x + 0.25, cy + 0.05, w - 0.5, h - (cy - y) - 0.2, body, size=body_size)


def lesson(slide, text, y=6.25):
    """배운 문장 한 줄 — 코랄 강조 카드."""
    rect(slide, 0.6, y, 12.1, 0.62, T["card2"] if THEME != "B" else "1E293B", MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
    tb(slide, 0.9, y + 0.13, 11.6, 0.4, [[("배운 문장  ", {"bold": True, "color": T["accent"], "size": 14}), (text, {"size": 14, "bold": True})]],
       size=14, anchor=MSO_ANCHOR.MIDDLE)


def table(slide, x, y, w, rows, col_widths, font_size=11, header_fill=None, row_h=0.36, first_col_bold=False, zebra=True):
    nrows, ncols = len(rows), len(rows[0])
    shp = slide.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tbl = shp.table
    tblPr = tbl._tbl.tblPr; tblPr.set("bandRow", "0"); tblPr.set("firstRow", "0")
    # 기본 스타일 제거
    style = tblPr.find(qn("a:tableStyleId"))
    if style is not None: tblPr.remove(style)
    for ci, cw in enumerate(col_widths):
        tbl.columns[ci].width = Inches(cw)
    for ri, row in enumerate(rows):
        tbl.rows[ri].height = Inches(row_h)
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            cell.margin_left = cell.margin_right = Inches(0.08); cell.margin_top = cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if ri == 0:
                cell.fill.fore_color.rgb = rgb(header_fill or T["primary"])
            else:
                cell.fill.fore_color.rgb = rgb(T["card"] if (zebra and ri % 2 == 0) else (T["bg"] if THEME != "B" else "172033"))
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                ln = etree.SubElement(tcPr, qn(tag)); ln.set("w", "6350"); ln.set("cap", "flat"); ln.set("cmpd", "sng")
                sf = etree.SubElement(ln, qn("a:solidFill")); c = etree.SubElement(sf, qn("a:srgbClr")); c.set("val", "FFFFFF" if ri == 0 else ("D9DCE3" if THEME != "B" else "334155"))
            tf = cell.text_frame; tf.word_wrap = True
            runs = val if isinstance(val, list) else [(str(val), {})]
            p = tf.paragraphs[0]
            for rtext, opts in runs:
                r = p.add_run(); r.text = rtext; f = r.font; f.name = FONT; f.size = Pt(opts.get("size", font_size))
                f.bold = opts.get("bold", ri == 0 or (first_col_bold and ci == 0))
                f.color.rgb = rgb(opts.get("color", ("FFFFFF" if ri == 0 else T["text"])))
    return shp


def mono(slide, x, y, w, h, text, size=10.5, fill=None):
    rect(slide, x, y, w, h, fill or T["card"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.03)
    lines = text.split("\n")
    box = slide.shapes.add_textbox(Inches(x + 0.2), Inches(y + 0.15), Inches(w - 0.4), Inches(h - 0.3))
    tf = box.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, l in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.line_spacing = 1.05
        r = p.add_run(); r.text = l if l else " "
        r.font.name = T["mono"]; r.font.size = Pt(size); r.font.color.rgb = rgb(T["text"])
    return box


def callout(slide, x, y, w, text, color=None, size=11):
    """작은 말풍선 캡션 (카드 옆 화살표 설명)."""
    rect(slide, x, y, w, 0.42, color or T["accent"], MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.4)
    tb(slide, x + 0.15, y + 0.08, w - 0.3, 0.3, text, size=size, bold=True, color="FFFFFF", anchor=MSO_ANCHOR.MIDDLE)


def trial_slide(num, title, lesson_text, cols, sub=None):
    """시행착오 슬라이드 공통 틀: 4열(증상/원인/알아챈 방법/바꾼 것) + 배운 문장."""
    s = prs.slides.add_slide(BLANK)
    header(s, num, title, sub)
    labels = ["증상", "원인", "어떻게 알아챘나", "바꾼 것"]
    icons = ["!", "?", "👁", "→"]
    colors = [T["accent"], T["accent"], T["primary"], T["good"]]
    cw, gap, x0, y0 = 2.85, 0.23, 0.6, 1.85
    for i, (lab, body) in enumerate(zip(labels, cols)):
        x = x0 + i * (cw + gap)
        card(s, x, y0, cw, 4.0, title=lab, body=body, title_color=colors[i], body_size=11.5, icon=icons[i] if icons[i] != "👁" else "◉")
    lesson(s, lesson_text, y=6.15)
    s.notes_slide.notes_text_frame.text = f"[{title}] 증상 → 원인 → 어떻게 알아챘나 → 바꾼 것 순서로 1분. 마지막에 배운 문장 한 줄을 그대로 읽는다: {lesson_text}"
    return s


