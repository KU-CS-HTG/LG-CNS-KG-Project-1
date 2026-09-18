"""팀원 덱(presentation_version2.pptx)의 우상단 'Y' 모티프를 우리 그래프 스키마 모양으로 교체.

    전공 ──▶ 역량 ◀── 직무      (2홉: DEVELOPS / REQUIRES)
              ▲
            하위 역량            (IS_A 한 단계, 위로만)

사용: motif_patch.py in.pptx out.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from lxml import etree

NAVY, CORAL, GREEN, GRAY, WHITE = "1E2761", "E8574A", "2F9E63", "9AA3B2", "FFFFFF"
FONT = "Malgun Gothic"


def rgb(h):
    return RGBColor.from_string(h)


def dot(slide, cx, cy, d, color, name):
    o = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    o.name = name
    o.fill.solid(); o.fill.fore_color.rgb = rgb(color); o.line.fill.background()
    o.shadow.inherit = False   # 빈 effectLst 를 넣어 테마 그림자 제거
    return o


def arrow(slide, x1, y1, x2, y2, color, width, name):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.name = name
    c.line.color.rgb = rgb(color); c.line.width = Pt(width)
    ln = c.line._get_or_add_ln()
    tail = etree.SubElement(ln, qn("a:tailEnd")); tail.set("type", "triangle"); tail.set("w", "sm"); tail.set("len", "sm")
    return c


def label(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.CENTER):
    t = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    t.name = "graph-label"
    tf = t.text_frame; tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.name = FONT; r.font.color.rgb = rgb(color)
    return t


def schema(slide, x0, y0, d, gap, line_w, labels=None, label_color=None, major_color=NAVY):
    """x0,y0 = 전공 노드 중심. d = 노드 지름, gap = 노드 중심 간격. labels = ('전공','역량','직무','하위 역량') or None."""
    xm, xs, xj = x0, x0 + gap, x0 + 2 * gap
    ys = y0 + gap * 0.8                    # 하위 역량은 역량 아래
    r = d / 2 + 0.02                        # 화살표는 노드 가장자리에서 시작/끝
    # 엣지: 전공 → 역량 (DEVELOPS), 직무 → 역량 (REQUIRES), 하위 → 역량 (IS_A)
    arrow(slide, xm + r, y0, xs - r, y0, major_color, line_w, "graph-edge-develops")
    arrow(slide, xj - r, y0, xs + r, y0, GREEN, line_w, "graph-edge-requires")
    arrow(slide, xs, ys - r, xs, y0 + r, GRAY, line_w, "graph-edge-isa")
    # 노드
    dot(slide, xm, y0, d, major_color, "graph-node-major")
    dot(slide, xs, y0, d, CORAL, "graph-node-skill")
    dot(slide, xj, y0, d, GREEN, "graph-node-job")
    dot(slide, xs, ys, d * 0.7, GRAY, "graph-node-skill-child")
    if labels:
        lw = gap * 0.95; ly = y0 + d / 2 + 0.05
        label(slide, xm - lw / 2, ly, lw, 0.25, labels[0], 9, label_color, bold=True)
        label(slide, xj - lw / 2, ly, lw, 0.25, labels[2], 9, label_color, bold=True)
        label(slide, xs - lw / 2, y0 - d / 2 - 0.24, lw, 0.25, labels[1], 9, label_color, bold=True)
        label(slide, xs + d * 0.35 + 0.03, ys - 0.1, lw, 0.25, labels[3], 8, label_color, align=PP_ALIGN.LEFT)


def strip_old_motif(slide):
    """우상단(x>11.9in, y<1.3in)에 있는 작은 도형(선·타원·연결선)만 제거. 제목·페이지 번호는 건드리지 않는다."""
    removed = []
    for sh in list(slide.shapes):
        if sh.left is None:
            continue
        x, y, w, h = sh.left / 914400, sh.top / 914400, sh.width / 914400, sh.height / 914400
        if x > 11.9 and y < 1.3 and w < 0.6 and h < 0.6 and not (sh.has_text_frame and sh.text_frame.text.strip()):
            removed.append(sh.name); sh._element.getparent().remove(sh._element)
    return removed


def main(src, dst):
    prs = Presentation(src)
    for i, s in enumerate(prs.slides, 1):
        gone = strip_old_motif(s)
        if i == 1:   # 표지: 크게 + 이름표 (흰 글씨, 짙은 배경)
            schema(s, 10.75, 0.85, 0.2, 0.85, 1.5, labels=("전공", "역량", "직무", "하위 역량"), label_color="D7DBEA", major_color="CADCFC")
        else:        # 본문: 작게, 제목 오른쪽 여백
            schema(s, 11.95, 0.58, 0.12, 0.4, 1.0)
        print(f"slide {i}: removed {len(gone)} -> {gone}")
    prs.save(dst); print("saved", dst)


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
